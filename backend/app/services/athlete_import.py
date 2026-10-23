import csv
import io
import unicodedata

from fastapi import HTTPException
from openpyxl import load_workbook

_REQUIRED_COLUMNS = ("NOME", "EQUIPE", "NAIPE", "CATEGORIA")
# REGISTRO e opcional para manter compatibilidade com planilhas antigas.
_IMPORT_COLUMNS = (*_REQUIRED_COLUMNS, "REGISTRO")


def _strip_accents(value: str) -> str:
    decomposed = unicodedata.normalize("NFKD", value)
    return "".join(
        char for char in decomposed if not unicodedata.combining(char)
    )


def _normalize_header(value: object) -> str:
    if value is None:
        return ""
    return _strip_accents(str(value)).strip().upper()


def _parse_csv_rows(data: bytes) -> list[list[object]]:
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = data.decode("latin-1")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    reader = csv.reader(io.StringIO(text), dialect)
    return [row for row in reader]


def _parse_xlsx_rows(
    data: bytes, max_rows: int | None = None
) -> list[list[object]]:
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        sheet = workbook.active
        if sheet is None:
            return []
        rows: list[list[object]] = []
        for row in sheet.iter_rows(values_only=True):
            rows.append(list(row))
            # Margem para cabecalho em qualquer uma das 20 primeiras linhas.
            if max_rows is not None and len(rows) > max_rows + 20:
                raise HTTPException(
                    status_code=413, detail="planilha_muito_grande"
                )
        return rows
    finally:
        workbook.close()


def _detect_header(rows: list[list[object]]) -> tuple[int, dict[str, int]]:
    for index, row in enumerate(rows[:20]):
        headers = [_normalize_header(cell) for cell in row]
        if "NOME" in headers:
            return index, {
                header: position
                for position, header in enumerate(headers)
                if header
            }
    raise HTTPException(status_code=400, detail="planilha_sem_cabecalho")


def _cell(row: list[object], position: int | None) -> str:
    if position is None or position >= len(row) or row[position] is None:
        return ""
    return str(row[position]).strip()
